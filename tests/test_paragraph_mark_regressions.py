from types import SimpleNamespace

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docxtor import project_docx_for_review

from reviewkit import (
    ReviewArtifactPreservationError,
    accept_all_revisions,
    assert_docx_structure_preserved,
    assess_rendered_actions,
)
from reviewkit.parser_docx import load_docx


def _marked_document(path, *, deleted_text=False, empty=False):
    doc = Document()
    paragraph = doc.add_paragraph(style="Heading 1")
    properties = OxmlElement("w:rPr")
    mark = OxmlElement("w:del")
    mark.set(qn("w:id"), "1")
    properties.append(mark)
    paragraph._p.get_or_add_pPr().append(properties)
    if deleted_text:
        deletion = OxmlElement("w:del")
        deletion.set(qn("w:id"), "2")
        run = OxmlElement("w:r")
        text = OxmlElement("w:delText")
        text.text = "Deleted heading."
        run.append(text)
        deletion.append(run)
        paragraph._p.append(deletion)
    elif not empty:
        paragraph.add_run("Heading.")
    doc.add_paragraph("Retained body.", style="Title")
    doc.save(path)


def test_paragraph_mark_only_revision_has_typed_coverage(tmp_path):
    path = tmp_path / "source.docx"
    _marked_document(path, empty=True)
    document = load_docx(path)
    entries = document.revision_ledger.entries
    assert len(entries) == 1
    assert entries[0].paragraph_mark is True
    assert entries[0].locator == "body:p:0"
    assert entries[0].text == ""
    assert document.revision_ledger.coverage.value == "complete"


def test_nonempty_deleted_boundary_fails_closed_until_origin_mapping_exists(tmp_path):
    path = tmp_path / "source.docx"
    _marked_document(path)
    document = load_docx(path)
    assert document.revision_ledger.coverage.value == "incomplete"


@pytest.mark.parametrize("tracked", [False, True])
def test_legacy_projection_without_boundary_api_does_not_claim_revision_coverage(
    tmp_path, monkeypatch, tracked
):
    path = tmp_path / "legacy.docx"
    if tracked:
        _marked_document(path, empty=True)
    else:
        doc = Document()
        doc.add_paragraph("Retained.")
        doc.save(path)
    projection = project_docx_for_review(path)
    legacy = SimpleNamespace(
        **{
            key: value
            for key, value in vars(projection).items()
            if key != "paragraph_mark_revisions"
        }
    )
    monkeypatch.setattr("reviewkit.parser_docx.project_docx_for_review", lambda _: legacy)
    document = load_docx(path)
    assert document.revision_ledger.coverage.value == ("incomplete" if tracked else "complete")


def test_corrected_preservation_allows_explicit_paragraph_mark_deletion(tmp_path):
    source, accepted = tmp_path / "source.docx", tmp_path / "accepted.docx"
    _marked_document(source, deleted_text=True)
    accept_all_revisions(source, accepted)
    assert_docx_structure_preserved(source, accepted, phase="corrected")
    assert [p.text for p in Document(accepted).paragraphs] == ["Retained body."]


def test_corrected_preservation_still_rejects_loss_of_retained_style(tmp_path):
    source, accepted = tmp_path / "source.docx", tmp_path / "accepted.docx"
    _marked_document(source, deleted_text=True)
    accept_all_revisions(source, accepted)
    doc = Document(accepted)
    doc.paragraphs[0]._p.get_or_add_pPr().remove(doc.paragraphs[0]._p.pPr.pStyle)
    doc.save(accepted)
    with pytest.raises(ReviewArtifactPreservationError, match="pStyle"):
        assert_docx_structure_preserved(source, accepted, phase="corrected")


def test_structural_boundary_receipts_are_not_unmatched_text_writing_actions(tmp_path):
    path = tmp_path / "source.docx"
    _marked_document(path)
    assessment = assess_rendered_actions(path, ())
    assert assessment.matches
    assert assessment.rendered_revision_count == 0
