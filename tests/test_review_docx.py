"""One DOCX is reviewed by walking that same file."""

from __future__ import annotations

from pathlib import Path
from zipfile import ZipFile

from docx import Document as DocxDocument

from reviewkit import (
    AKAPIT,
    CALOSC,
    ROZDZIAL,
    ZDANIE,
    StayOrGo,
    load_docx,
    read_comments,
    review_docx,
)
from reviewkit.live_docx import LiveDocx
from reviewkit.walk import ReviewUnit

_FIRST = "First sentence."
_SECOND = "Second sentence."
_THIRD = "Third sentence lives here."
_THIRD_EDITED = "Third sentence was edited."
_FOURTH = "Fourth sentence to drop."
_INSERTED = " Extra sentence."
_ALPHA = "Alpha sentence."
_BRAVO = "Bravo sentence."
_CHARLIE = "Charlie sentence."
_DELTA = "Delta sentence."
_NOTE_ONE = "first note on the opening sentence"
_NOTE_TWO = "second note on the same sentence"
_NOTE_UPDATED = "updated note on the opening sentence"


def _sample_docx(path: Path) -> Path:
    document = DocxDocument()
    document.add_heading("Rozdział 1", level=1)
    document.add_paragraph(f"{_FIRST} {_SECOND}")
    document.add_paragraph(_THIRD)
    document.add_paragraph(_FOURTH)
    document.save(path)
    return path


def _document_xml(path: Path) -> bytes:
    with ZipFile(path) as archive:
        return archive.read("word/document.xml")


def _three_sentence_docx(path: Path) -> Path:
    document = DocxDocument()
    document.add_heading("Rozdział 1", level=1)
    document.add_paragraph(f"{_ALPHA} {_BRAVO} {_CHARLIE}")
    document.save(path)
    return path


def _delete_unit_text(docx: LiveDocx, unit: ReviewUnit) -> None:
    assert unit.paragraph_locator is not None
    assert unit.char_start is not None
    assert unit.char_end is not None
    docx.delete_text(
        locator=unit.paragraph_locator,
        start=unit.char_start,
        end=unit.char_end,
        expected_text=unit.text,
    )


class ScriptedReviewer:
    def __init__(self) -> None:
        self.comments_before_second_sentence: list[str] | None = None
        self.second_sentence_seen = False
        self.comment_ids: list[str] = []
        self.saw_update_on_file = False
        self.saw_insert_on_file = False
        self.saw_delete_on_file = False
        self.saw_replace_on_file = False
        self.calosc_stays = 0

    def decide(self, unit: ReviewUnit, docx: LiveDocx) -> StayOrGo:
        if unit.level == ZDANIE and unit.text == _FIRST and unit.stay_index == 0:
            assert self.second_sentence_seen is False
            self.comment_ids.append(_comment_on_unit(docx, unit, _NOTE_ONE))
            return StayOrGo.STAY
        if unit.level == ZDANIE and unit.text == _FIRST and unit.stay_index == 1:
            assert self.second_sentence_seen is False
            self.comment_ids.append(_comment_on_unit(docx, unit, _NOTE_TWO))
            sitting = {comment.text for comment in read_comments(docx.path)}
            assert sitting >= {_NOTE_ONE, _NOTE_TWO}
            self.comments_before_second_sentence = [
                comment.text for comment in read_comments(docx.path)
            ]
            return StayOrGo.STAY
        if unit.level == ZDANIE and unit.text == _FIRST and unit.stay_index == 2:
            assert self.second_sentence_seen is False
            return StayOrGo.GO
        if unit.level == ZDANIE and unit.text == _SECOND:
            self.second_sentence_seen = True
            assert self.comments_before_second_sentence is not None
            return StayOrGo.GO
        if unit.level == ZDANIE and _THIRD in unit.text and not self.saw_replace_on_file:
            assert unit.paragraph_locator is not None
            assert unit.char_start is not None
            assert unit.char_end is not None
            docx.replace_text(
                locator=unit.paragraph_locator,
                start=unit.char_start,
                end=unit.char_end,
                replacement=_THIRD_EDITED,
                expected_text=unit.text,
            )
            assert _THIRD_EDITED in docx.load().text
            self.saw_replace_on_file = True
            return StayOrGo.GO
        if unit.level == ZDANIE and unit.text == _FOURTH:
            assert unit.paragraph_locator is not None
            assert unit.char_start is not None
            assert unit.char_end is not None
            docx.delete_text(
                locator=unit.paragraph_locator,
                start=unit.char_start,
                end=unit.char_end,
                expected_text=unit.text,
            )
            self.saw_delete_on_file = True
            return StayOrGo.GO
        if unit.level == AKAPIT and _FIRST in unit.text and not self.saw_insert_on_file:
            assert unit.paragraph_locator is not None
            docx.insert_text(
                locator=unit.paragraph_locator,
                offset=len(unit.text),
                text=_INSERTED,
            )
            assert _INSERTED.strip() in docx.load().text
            self.saw_insert_on_file = True
            return StayOrGo.GO
        if unit.level == ROZDZIAL and not self.saw_update_on_file:
            self.comment_ids[0] = docx.update_comment(self.comment_ids[0], _NOTE_UPDATED)
            texts = [comment.text for comment in read_comments(docx.path)]
            assert _NOTE_UPDATED in texts
            assert _NOTE_ONE not in texts
            remaining = next(
                comment.id for comment in read_comments(docx.path) if comment.text == _NOTE_TWO
            )
            docx.delete_comment(remaining)
            texts = [comment.text for comment in read_comments(docx.path)]
            assert _NOTE_TWO not in texts
            self.saw_update_on_file = True
            return StayOrGo.GO
        if unit.level == CALOSC:
            self.calosc_stays += 1
            if unit.stay_index == 0:
                return StayOrGo.STAY
            return StayOrGo.GO
        return StayOrGo.GO


def _comment_on_unit(docx: LiveDocx, unit: ReviewUnit, text: str) -> str:
    assert unit.paragraph_locator is not None
    assert unit.char_start is not None
    assert unit.char_end is not None
    return docx.add_comment(
        locator=unit.paragraph_locator,
        start=unit.char_start,
        end=unit.char_end,
        text=text,
        expected_text=unit.text,
    )


def test_live_docx_keeps_one_open_handle(tmp_path: Path) -> None:
    path = _sample_docx(tmp_path / "handle.docx")
    live = LiveDocx(path)
    handle = live.open()
    comment_id = live.add_comment(
        locator="body:p:1",
        start=0,
        end=len(_FIRST),
        text=_NOTE_ONE,
        expected_text=_FIRST,
    )
    assert live.open() is handle
    live.update_comment(comment_id, _NOTE_UPDATED)
    assert live.open() is handle
    live.delete_comment(comment_id)
    assert live.open() is handle
    live.insert_text(locator="body:p:1", offset=len(f"{_FIRST} {_SECOND}"), text=_INSERTED)
    assert live.open() is handle


def test_review_walks_one_docx_in_place(tmp_path: Path) -> None:
    path = _sample_docx(tmp_path / "source.docx")
    reviewer = ScriptedReviewer()

    result = review_docx(path, reviewer)

    assert result.path == path
    levels = [visit.level for visit in result.visits]
    assert levels[0] == ZDANIE
    assert ROZDZIAL not in levels[: levels.index(AKAPIT)]
    assert CALOSC not in levels[: levels.index(ROZDZIAL)]
    zdanie = [visit for visit in result.visits if visit.level == ZDANIE]
    assert all(visit.level == ZDANIE for visit in result.visits[: len(zdanie)])
    assert [visit.stay_index for visit in zdanie[:3]] == [0, 1, 2]
    assert zdanie[0].text == _FIRST
    assert zdanie[1].text == _FIRST
    assert zdanie[2].text == _FIRST
    assert zdanie[3].text == _SECOND
    assert reviewer.second_sentence_seen is True
    assert reviewer.comments_before_second_sentence is not None
    assert reviewer.comments_before_second_sentence.count(_NOTE_ONE) == 1
    assert reviewer.comments_before_second_sentence.count(_NOTE_TWO) == 1

    akapit_at = levels.index(AKAPIT)
    rozdzial_at = levels.index(ROZDZIAL)
    calosc_at = levels.index(CALOSC)
    assert akapit_at < rozdzial_at < calosc_at
    assert levels.count(AKAPIT) >= 2
    assert ROZDZIAL in levels
    calosc = [visit for visit in result.visits if visit.level == CALOSC]
    assert [visit.stay_index for visit in calosc] == [0, 1]
    assert reviewer.calosc_stays == 2

    assert reviewer.saw_update_on_file
    assert reviewer.saw_insert_on_file
    assert reviewer.saw_delete_on_file
    assert reviewer.saw_replace_on_file
    texts = [comment.text for comment in read_comments(path)]
    assert _NOTE_UPDATED in texts
    assert _NOTE_TWO not in texts
    assert _NOTE_ONE not in texts
    document = load_docx(path)
    assert _THIRD_EDITED in document.text
    assert _INSERTED.strip() in document.text
    assert document.source_path == path
    xml = _document_xml(path)
    assert b"w:ins" in xml
    assert b"w:del" in xml


def test_delete_middle_sentence_then_dalej_continues_walk(tmp_path: Path) -> None:
    path = _three_sentence_docx(tmp_path / "delete-dalej.docx")

    def decide(unit: ReviewUnit, docx: LiveDocx) -> StayOrGo:
        if unit.level == ZDANIE and unit.text == _BRAVO:
            _delete_unit_text(docx, unit)
        return StayOrGo.GO

    result = review_docx(path, decide)

    zdanie = [visit for visit in result.visits if visit.level == ZDANIE]
    assert [visit.text for visit in zdanie] == [_ALPHA, _BRAVO, _CHARLIE]
    levels = [visit.level for visit in result.visits]
    assert AKAPIT in levels
    assert ROZDZIAL in levels
    assert CALOSC in levels
    assert levels.index(ZDANIE) < levels.index(AKAPIT) < levels.index(ROZDZIAL) < levels.index(CALOSC)


def test_delete_then_zostan_does_not_continue_on_next_sentence(tmp_path: Path) -> None:
    path = _three_sentence_docx(tmp_path / "delete-zostan.docx")

    def decide(unit: ReviewUnit, docx: LiveDocx) -> StayOrGo:
        if unit.level == ZDANIE and unit.text == _BRAVO and unit.stay_index == 0:
            _delete_unit_text(docx, unit)
            return StayOrGo.STAY
        if unit.level == ZDANIE and unit.stay_index > 0:
            raise AssertionError(f"stay continued onto {unit.text!r}")
        return StayOrGo.GO

    result = review_docx(path, decide)

    zdanie = [visit for visit in result.visits if visit.level == ZDANIE]
    assert [(visit.text, visit.stay_index) for visit in zdanie] == [
        (_ALPHA, 0),
        (_BRAVO, 0),
        (_CHARLIE, 0),
    ]
    levels = [visit.level for visit in result.visits]
    assert AKAPIT in levels
    assert ROZDZIAL in levels
    assert CALOSC in levels


def test_new_sentence_created_during_zdanie_is_visited(tmp_path: Path) -> None:
    path = _three_sentence_docx(tmp_path / "insert-zdanie.docx")
    inserted = False

    def decide(unit: ReviewUnit, docx: LiveDocx) -> StayOrGo:
        nonlocal inserted
        if unit.level == ZDANIE and unit.text == _ALPHA and not inserted:
            assert unit.paragraph_locator is not None
            paragraph = next(
                node
                for node in docx.load().iter_paragraphs()
                if node.locator == unit.paragraph_locator
            )
            docx.insert_text(
                locator=unit.paragraph_locator,
                offset=len(paragraph.text),
                text=f" {_DELTA}",
            )
            inserted = True
        return StayOrGo.GO

    result = review_docx(path, decide)

    zdanie = [visit.text for visit in result.visits if visit.level == ZDANIE]
    assert zdanie == [_ALPHA, _BRAVO, _CHARLIE, _DELTA]
    levels = [visit.level for visit in result.visits]
    assert AKAPIT in levels
    assert ROZDZIAL in levels
    assert CALOSC in levels


def test_readme_describes_only_the_stay_or_go_walk() -> None:
    readme = (Path(__file__).resolve().parents[1] / "README.md").read_text(encoding="utf-8")
    assert "```mermaid" in readme
    assert readme.count("```mermaid") == 1
    mermaid = readme.split("```mermaid", 1)[1].split("```", 1)[0]
    for name in ("zdanie", "akapit", "rozdział", "całość"):
        assert name in mermaid
    assert "zdanie 1" in mermaid
    assert "zdanie 2" in mermaid
    assert "akapit 1" in mermaid
    assert "akapit 2" in mermaid
    assert "rozdział 1" in mermaid
    assert "rozdział 2" in mermaid
    assert "stay or go" in mermaid
    assert mermaid.count("stay or go") >= 4
    assert "zNext --> zLoop" not in mermaid
    assert "comment or text on this DOCX" in mermaid
    assert "walk całość again, n times" in mermaid
    assert "same DOCX" in mermaid
    walk = readme.split("## The walk")[1].split("##")[0]
    assert "grain" not in walk
    assert "percent" not in walk
    assert "sides" not in walk
    lowered = readme.lower()
    assert "review_tree" not in lowered
    assert "review_document" not in lowered
    assert "two-scan" not in lowered
    assert "reviewed.docx" not in lowered
