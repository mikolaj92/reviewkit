from __future__ import annotations

import hashlib
import json
from dataclasses import replace
from pathlib import Path
from typing import Any

import pytest
from docx import Document
from docx.comments import Comment
from docx.text.run import Run
from docxtor import ComparisonDiagnostic, compare_docx_documents

import reviewkit
from reviewkit import (
    ActionStatus,
    EvidenceRef,
    FindingLineageEvent,
    ReviewAction,
    ReviewActionType,
    ReviewReference,
    ReviewScope,
    accept_all_revisions,
    load_docx,
)
from reviewkit.comment_formatter import format_action_comment, format_legacy_action_comment
from reviewkit.renderer_docx import render_reviewed_docx

_ACTION_SHA256 = "b65fb2162f332fe521617bd3c040badaa15f1a2ee24e641db2cdc321465136be"
_LEGACY_ACTION_SHA256 = "a1e4d4ee3b01329714911d1b0beb18d951470d455e983489148a013e45c913bf"
_SOURCE_COMMENT = "Status: applied — komentarz obecny przed procesem."
_ACTION_COMMENT = "Dostosowano termin zapłaty."
_ACTION_ID = "term-30"
_HISTORICAL_ACTION_COMMENT = (
    "CORRECTION: Replace marker.\n"
    "Original: 'alpha'\n"
    "Replacement: 'beta'\n"
    "Category: term\n"
    "Policy: policy check\n"
    "References: Section 1\n"
    "Evidence: body:p:0\n"
    "Status: applied"
)


def test_verified_action_provenance_covers_all_three_document_pairs(tmp_path: Path) -> None:
    source, reviewed, corrected, action = _review_fixture(tmp_path)
    evidence = _evidence(source, reviewed, [action])
    transition = _transition(reviewed, corrected)
    attribute = _public_api("attribute_docx_changes")

    input_reviewed = compare_docx_documents(source, reviewed)
    input_reviewed_provenance = attribute(
        input_reviewed,
        action_source=source,
        review_evidence=evidence,
    )

    text = _one(input_reviewed_provenance, "text")
    assert text.status == "verified_action"
    assert text.action_ids == (_ACTION_ID,)
    revisions = _events(input_reviewed_provenance, "revision")
    assert len(revisions) == 2
    assert {event.status for event in revisions} == {"verified_action"}
    assert {event.action_ids for event in revisions} == {(_ACTION_ID,)}
    comments = _events(input_reviewed_provenance, "comment")
    assert len(comments) == 2
    assert sum(event.status == "verified_action" for event in comments) == 1

    source_comment = next(
        event
        for event in input_reviewed.comment_changes
        if event.left is not None and event.left.text == _SOURCE_COMMENT
    )
    source_comment_result = _by_change_id(input_reviewed_provenance, source_comment.change_id)
    assert source_comment_result.status == "source"
    assert source_comment_result.action_ids == ()

    reviewed_corrected = compare_docx_documents(reviewed, corrected)
    reviewed_corrected_provenance = attribute(
        reviewed_corrected,
        action_source=source,
        review_evidence=evidence,
        transition_evidence=transition,
        source_to_reviewed=input_reviewed,
    )
    assert _events(reviewed_corrected_provenance, "text") == []
    accepted_revisions = _events(reviewed_corrected_provenance, "revision")
    assert len(accepted_revisions) == 2
    assert {event.status for event in accepted_revisions} == {"verified_acceptance"}
    assert {event.action_ids for event in accepted_revisions} == {(_ACTION_ID,)}
    removed_comments = _events(reviewed_corrected_provenance, "comment")
    assert len(removed_comments) == 2
    assert sum(event.status == "source" for event in removed_comments) == 1
    assert sum(event.status == "verified_acceptance" for event in removed_comments) == 1

    input_corrected = compare_docx_documents(source, corrected)
    input_corrected_provenance = attribute(
        input_corrected,
        action_source=source,
        review_evidence=evidence,
        transition_evidence=transition,
    )
    corrected_text = _one(input_corrected_provenance, "text")
    assert corrected_text.status == "verified_action"
    assert corrected_text.action_ids == (_ACTION_ID,)
    assert _events(input_corrected_provenance, "revision") == []
    removed_source_comment = next(
        event
        for event in input_corrected.comment_changes
        if event.left is not None and event.left.text == _SOURCE_COMMENT
    )
    assert (
        _by_change_id(input_corrected_provenance, removed_source_comment.change_id).status
        == "source"
    )


@pytest.mark.parametrize(
    "with_lineage",
    [False, True],
    ids=["pre-lineage", "lineage"],
)
def test_historical_comment_projection_is_exact_for_hash_bound_evidence(
    tmp_path: Path, with_lineage: bool
) -> None:
    source, reviewed, corrected, _action, evidence = _historical_comment_fixture(
        tmp_path,
        with_lineage=with_lineage,
    )
    input_reviewed = compare_docx_documents(source, reviewed)
    direct = _public_api("attribute_docx_changes")(
        input_reviewed,
        action_source=source,
        review_evidence=evidence,
    )
    action_comments = [
        event for event in _events(direct, "comment") if event.status == "verified_action"
    ]
    assert len(action_comments) == 1
    assert action_comments[0].action_ids == ("replace-marker",)

    transition = _transition(reviewed, corrected)
    accepted = _public_api("attribute_docx_changes")(
        compare_docx_documents(reviewed, corrected),
        action_source=source,
        review_evidence=evidence,
        transition_evidence=transition,
        source_to_reviewed=input_reviewed,
    )
    accepted_comments = [
        event for event in _events(accepted, "comment") if event.status == "verified_acceptance"
    ]
    assert len(accepted_comments) == 1
    assert accepted_comments[0].action_ids == ("replace-marker",)


def test_historical_comment_payload_mismatch_stays_unknown(tmp_path: Path) -> None:
    source, reviewed, _corrected, _action, evidence = _historical_comment_fixture(
        tmp_path,
        with_lineage=True,
        comment_text=_HISTORICAL_ACTION_COMMENT.replace("Status: applied", "Status: applied!"),
    )
    comparison = compare_docx_documents(source, reviewed)
    provenance = _public_api("attribute_docx_changes")(
        comparison,
        action_source=source,
        review_evidence=evidence,
    )

    action_comment = next(event for event in comparison.comment_changes if event.left is None)
    attributed = _by_change_id(provenance, action_comment.change_id)
    assert attributed.status == "unknown"
    assert attributed.action_ids == ()
    assert attributed.evidence_codes == ("no_exact_provenance_match",)


def test_historical_comment_reviewed_hash_mismatch_fails_closed(tmp_path: Path) -> None:
    source, reviewed, _corrected, _action, evidence = _historical_comment_fixture(
        tmp_path,
        with_lineage=True,
    )
    reviewed_document = Document(reviewed)
    reviewed_document.paragraphs[0].runs[-1].text += " tampered"
    reviewed_document.save(reviewed)

    comparison = compare_docx_documents(source, reviewed)
    provenance = _public_api("attribute_docx_changes")(
        comparison,
        action_source=source,
        review_evidence=evidence,
    )
    dependent = [*_events(provenance, "text"), *_events(provenance, "revision")]
    assert dependent
    assert {event.status for event in dependent} == {"unknown"}
    assert all(event.action_ids == () for event in dependent)
    assert all("comparison_hash_chain_mismatch" in event.evidence_codes for event in dependent)


def test_historical_comment_bad_anchor_stays_unknown(tmp_path: Path) -> None:
    source, reviewed, _corrected, action, _evidence = _historical_comment_fixture(
        tmp_path,
        with_lineage=True,
    )
    bad_action = action.model_copy(update={"node_id": "body:p:missing"})
    evidence = _evidence_for_actions(source, reviewed, [bad_action])
    comparison = compare_docx_documents(source, reviewed)
    provenance = _public_api("attribute_docx_changes")(
        comparison,
        action_source=source,
        review_evidence=evidence,
    )

    dependent = [*_events(provenance, "text"), *_events(provenance, "revision")]
    assert dependent
    assert {event.status for event in dependent} == {"unknown"}
    assert all(event.action_ids == () for event in dependent)
    action_comment = next(event for event in comparison.comment_changes if event.left is None)
    assert _by_change_id(provenance, action_comment.change_id).status == "unknown"


def test_historical_duplicate_same_anchor_comments_stay_unknown(tmp_path: Path) -> None:
    source, reviewed, _corrected, action, _evidence = _historical_comment_fixture(
        tmp_path,
        with_lineage=True,
    )
    reviewed_document = Document(reviewed)
    paragraph = reviewed_document.paragraphs[0]
    action_run = Run(paragraph._p.xpath("./w:ins/w:r")[0], paragraph)
    reviewed_document.add_comment(
        runs=[action_run],
        text=_HISTORICAL_ACTION_COMMENT,
        author="Reviewer",
    )
    reviewed_document.save(reviewed)
    evidence = _evidence_for_actions(source, reviewed, [action])

    comparison = compare_docx_documents(source, reviewed)
    provenance = _public_api("attribute_docx_changes")(
        comparison,
        action_source=source,
        review_evidence=evidence,
    )
    duplicate_events = [
        event
        for event in comparison.comment_changes
        if event.left is None and event.right is not None
    ]
    assert len(duplicate_events) == 2
    assert all(event.right.anchor is not None for event in duplicate_events)
    duplicate_anchors = {
        (
            event.right.anchor.locator,
            event.right.anchor.block_id,
            event.right.anchor.pair_id,
            event.right.anchor.start_offset,
            event.right.anchor.end_offset,
        )
        for event in duplicate_events
        if event.right.anchor is not None
    }
    assert len(duplicate_anchors) == 1
    action_comments = [_by_change_id(provenance, event.change_id) for event in duplicate_events]
    assert len(action_comments) == 2
    assert {event.status for event in action_comments} == {"unknown"}
    assert all(event.action_ids == () for event in action_comments)
    assert {event.evidence_codes for event in action_comments} == {("no_exact_provenance_match",)}


def test_dike_authored_source_comment_is_not_process_provenance(tmp_path: Path) -> None:
    source, reviewed, _corrected, action = _review_fixture(tmp_path)
    evidence = _evidence(source, reviewed, [action])
    comparison = compare_docx_documents(source, reviewed)
    provenance = _public_api("attribute_docx_changes")(
        comparison,
        action_source=source,
        review_evidence=evidence,
    )

    dike_authored = next(
        event
        for event in comparison.comment_changes
        if event.left is not None
        and event.left.author == "Dike"
        and event.left.text == _SOURCE_COMMENT
    )
    attributed = _by_change_id(provenance, dike_authored.change_id)
    assert attributed.status == "source"
    assert attributed.action_ids == ()


def test_broken_action_hash_chain_leaves_text_and_revisions_unknown(tmp_path: Path) -> None:
    source, reviewed, _corrected, action = _review_fixture(tmp_path)
    evidence = replace(_evidence(source, reviewed, [action]), actions_sha256="0" * 64)
    comparison = compare_docx_documents(source, reviewed)

    provenance = _public_api("attribute_docx_changes")(
        comparison,
        action_source=source,
        review_evidence=evidence,
    )

    dependent = [*_events(provenance, "text"), *_events(provenance, "revision")]
    assert dependent
    assert {event.status for event in dependent} == {"unknown"}
    assert all(event.action_ids == () for event in dependent)


def test_pre_lineage_action_sidecar_hash_is_verified_without_ignoring_lineage(
    tmp_path: Path,
) -> None:
    source, reviewed, _corrected, action = _review_fixture(tmp_path)
    legacy_payload = action.model_dump(mode="json", by_alias=True)
    legacy_payload.pop("lineage")
    legacy_digest = _json_digest([legacy_payload])
    assert legacy_digest == _LEGACY_ACTION_SHA256

    legacy_action = ReviewAction.model_validate(legacy_payload)
    assert "lineage" not in legacy_action.model_fields_set
    evidence_type = _public_api("ReviewActionEvidence")
    legacy_evidence = evidence_type(
        source_sha256=_sha256(source),
        reviewed_sha256=_sha256(reviewed),
        actions_sha256=legacy_digest,
        source_revision_signatures=reviewkit.revision_signatures(source),
        actions=(legacy_action,),
    )
    comparison = compare_docx_documents(source, reviewed)
    compatible = _public_api("attribute_docx_changes")(
        comparison,
        action_source=source,
        review_evidence=legacy_evidence,
    )
    assert _one(compatible, "text").status == "verified_action"

    current_payload = {**legacy_payload, "lineage": [{"kind": "review"}]}
    current_action = ReviewAction.model_validate(current_payload)
    assert "lineage" in current_action.model_fields_set
    mismatched_schema = _public_api("attribute_docx_changes")(
        comparison,
        action_source=source,
        review_evidence=replace(legacy_evidence, actions=(current_action,)),
    )
    dependent = [*_events(mismatched_schema, "text"), *_events(mismatched_schema, "revision")]
    assert dependent
    assert {event.status for event in dependent} == {"unknown"}
    assert all(event.action_ids == () for event in dependent)


def test_missing_or_mismatched_hash_evidence_fails_closed(tmp_path: Path) -> None:
    source, reviewed, corrected, action = _review_fixture(tmp_path)
    evidence = _evidence(source, reviewed, [action])
    input_reviewed = compare_docx_documents(source, reviewed)
    invalid_evidence = (
        None,
        replace(evidence, source_sha256="0" * 64),
        replace(evidence, reviewed_sha256="0" * 64),
        replace(evidence, source_revision_signatures=("not-the-source-signature",)),
    )

    for item in invalid_evidence:
        provenance = _public_api("attribute_docx_changes")(
            input_reviewed,
            action_source=source,
            review_evidence=item,
        )
        dependent = [*_events(provenance, "text"), *_events(provenance, "revision")]
        assert dependent
        assert {event.status for event in dependent} == {"unknown"}
        assert all(event.action_ids == () for event in dependent)

    input_corrected = compare_docx_documents(source, corrected)
    bad_transition = replace(_transition(reviewed, corrected), output_sha256="0" * 64)
    corrected_provenance = _public_api("attribute_docx_changes")(
        input_corrected,
        action_source=source,
        review_evidence=evidence,
        transition_evidence=bad_transition,
    )
    assert {event.status for event in _events(corrected_provenance, "text")} == {"unknown"}


def test_duplicate_action_ids_and_repeated_target_quotes_stay_unknown(tmp_path: Path) -> None:
    source, reviewed, _corrected, action = _review_fixture(tmp_path)
    comparison = compare_docx_documents(source, reviewed)
    duplicate_evidence_type = _public_api("ReviewActionEvidence")
    duplicate_actions = (action, action.model_copy())
    duplicate_evidence = duplicate_evidence_type(
        source_sha256=_sha256(source),
        reviewed_sha256=_sha256(reviewed),
        actions_sha256=_action_digest(list(duplicate_actions)),
        source_revision_signatures=(),
        actions=duplicate_actions,
    )
    duplicate_result = _public_api("attribute_docx_changes")(
        comparison,
        action_source=source,
        review_evidence=duplicate_evidence,
    )
    assert {event.status for event in _events(duplicate_result, "text")} == {"unknown"}
    assert {event.status for event in _events(duplicate_result, "revision")} == {"unknown"}

    repeated_source = tmp_path / "repeated-input.docx"
    repeated_reviewed = tmp_path / "repeated-reviewed.docx"
    repeated_document = Document()
    repeated_document.add_paragraph("The term is 14 days, then 14 more days.")
    repeated_document.save(repeated_source)
    parsed = load_docx(repeated_source)
    target = next(parsed.iter_paragraphs())
    repeated_action = ReviewAction(
        id="ambiguous-14",
        scope=ReviewScope.PARAGRAPH,
        action_type=ReviewActionType.REPLACE_TEXT,
        node_id=target.id,
        original_text="14",
        replacement_text="30",
        comment="Update the term.",
        status=ActionStatus.APPLIED,
        apply_to_corrected=True,
    )
    render_reviewed_docx(parsed, [repeated_action], repeated_reviewed)
    repeated_evidence = _evidence_for_actions(repeated_source, repeated_reviewed, [repeated_action])
    repeated_result = _public_api("attribute_docx_changes")(
        compare_docx_documents(repeated_source, repeated_reviewed),
        action_source=repeated_source,
        review_evidence=repeated_evidence,
    )
    dependent = [*_events(repeated_result, "text"), *_events(repeated_result, "revision")]
    assert dependent
    assert {event.status for event in dependent} == {"unknown"}
    assert all(event.action_ids == () for event in dependent)


def test_input_to_corrected_attributes_exact_insert_and_delete_actions(tmp_path: Path) -> None:
    source = tmp_path / "edits-input.docx"
    reviewed = tmp_path / "edits-reviewed.docx"
    corrected = tmp_path / "edits-corrected.docx"
    document = Document()
    document.add_paragraph("Remove obsolete wording from this clause.")
    document.add_paragraph("Payment is due on receipt.")
    document.save(source)

    parsed = load_docx(source)
    paragraphs = tuple(parsed.iter_paragraphs())
    delete_action = ReviewAction(
        id="delete-obsolete",
        scope=ReviewScope.PARAGRAPH,
        action_type=ReviewActionType.DELETE_TEXT,
        node_id=paragraphs[0].id,
        original_text="obsolete",
        status=ActionStatus.APPLIED,
        apply_to_corrected=True,
    )
    insert_action = ReviewAction(
        id="insert-confirmation",
        scope=ReviewScope.PARAGRAPH,
        action_type=ReviewActionType.INSERT_AFTER,
        node_id=paragraphs[1].id,
        original_text="Payment",
        replacement_text=" confirmed",
        status=ActionStatus.APPLIED,
        apply_to_corrected=True,
    )
    actions = [delete_action, insert_action]
    render_reviewed_docx(parsed, actions, reviewed)
    accept_all_revisions(reviewed, corrected)

    evidence = _evidence_for_actions(source, reviewed, actions)
    provenance = _public_api("attribute_docx_changes")(
        compare_docx_documents(source, corrected),
        action_source=source,
        review_evidence=evidence,
        transition_evidence=_transition(reviewed, corrected),
    )

    text_changes = _events(provenance, "text")
    assert {event.action_ids for event in text_changes if event.status == "verified_action"} == {
        ("delete-obsolete",),
        ("insert-confirmation",),
    }


def test_insertion_at_wrong_position_is_unknown_across_all_document_pairs(
    tmp_path: Path,
) -> None:
    source = tmp_path / "position-input.docx"
    reviewed = tmp_path / "position-reviewed.docx"
    corrected = tmp_path / "position-corrected.docx"
    document = Document()
    document.add_paragraph("Alpha beta gamma.")
    document.save(source)

    parsed = load_docx(source)
    target = next(parsed.iter_paragraphs())
    persisted_action = ReviewAction(
        id="insert-after-alpha",
        scope=ReviewScope.PARAGRAPH,
        action_type=ReviewActionType.INSERT_AFTER,
        node_id=target.id,
        original_text="Alpha",
        replacement_text=" EXTRA",
        comment="Insert the missing term.",
        status=ActionStatus.APPLIED,
        apply_to_corrected=True,
    )
    # The package contains the same payload in the same paragraph, but at beta.
    # A target quote and a unique insertion string alone must not prove position.
    rendered_action = persisted_action.model_copy(update={"original_text": "beta"})
    render_reviewed_docx(parsed, [rendered_action], reviewed)
    accept_all_revisions(reviewed, corrected)
    evidence = _evidence_for_actions(source, reviewed, [persisted_action])
    transition = _transition(reviewed, corrected)
    attribute = _public_api("attribute_docx_changes")

    input_reviewed = compare_docx_documents(source, reviewed)
    direct = attribute(
        input_reviewed,
        action_source=source,
        review_evidence=evidence,
    )
    assert _events(direct, "revision")
    assert {event.status for event in _events(direct, "revision")} == {"unknown"}
    assert {event.action_ids for event in _events(direct, "revision")} == {()}
    assert _events(direct, "text")
    assert {event.status for event in _events(direct, "text")} == {"unknown"}

    reviewed_corrected = attribute(
        compare_docx_documents(reviewed, corrected),
        action_source=source,
        review_evidence=evidence,
        transition_evidence=transition,
        source_to_reviewed=input_reviewed,
    )
    assert _events(reviewed_corrected, "revision")
    assert {event.status for event in _events(reviewed_corrected, "revision")} == {"unknown"}
    assert {event.action_ids for event in _events(reviewed_corrected, "revision")} == {()}

    input_corrected = attribute(
        compare_docx_documents(source, corrected),
        action_source=source,
        review_evidence=evidence,
        transition_evidence=transition,
    )
    assert _events(input_corrected, "text")
    assert {event.status for event in _events(input_corrected, "text")} == {"unknown"}
    assert {event.action_ids for event in _events(input_corrected, "text")} == {()}


def test_duplicate_same_anchor_action_comments_remain_unknown(tmp_path: Path) -> None:
    source = tmp_path / "duplicate-comment-input.docx"
    reviewed = tmp_path / "duplicate-comment-reviewed.docx"
    document = Document()
    document.add_paragraph("Alpha beta gamma.")
    document.save(source)

    parsed = load_docx(source)
    target = next(parsed.iter_paragraphs())
    action = ReviewAction(
        id="comment-alpha",
        scope=ReviewScope.PARAGRAPH,
        action_type=ReviewActionType.COMMENT,
        node_id=target.id,
        original_text="Alpha beta gamma.",
        comment="Check this clause.",
        status=ActionStatus.APPLIED,
        apply_to_corrected=True,
    )
    render_reviewed_docx(parsed, [action], reviewed)
    # Add a second byte-equivalent generated comment at the same selection.
    reviewed_document = Document(reviewed)
    reviewed_document.add_comment(
        runs=reviewed_document.paragraphs[0].runs,
        text=format_action_comment(action),
        author="Dike",
    )
    reviewed_document.save(reviewed)

    evidence = _evidence_for_actions(source, reviewed, [action])
    comparison = compare_docx_documents(source, reviewed)
    provenance = _public_api("attribute_docx_changes")(
        comparison,
        action_source=source,
        review_evidence=evidence,
    )

    comments = _events(provenance, "comment")
    assert len(comments) == 2
    assert {event.status for event in comments} == {"unknown"}
    assert {event.action_ids for event in comments} == {()}


def test_unrelated_numbering_preview_gap_keeps_exact_text_attribution(tmp_path: Path) -> None:
    source = tmp_path / "numbering-input.docx"
    reviewed = tmp_path / "numbering-reviewed.docx"
    document = Document()
    document.add_paragraph("The payment term is 14 days.")
    document.add_paragraph("Unrelated numbered paragraph.", style="List Number")
    document.save(source)

    parsed = load_docx(source)
    target = next(item for item in parsed.iter_paragraphs() if "14 days" in item.text)
    action = ReviewAction(
        id="replace-payment-term",
        scope=ReviewScope.PARAGRAPH,
        action_type=ReviewActionType.REPLACE_TEXT,
        node_id=target.id,
        original_text="14",
        replacement_text="30",
        comment="Update term.",
        status=ActionStatus.APPLIED,
        apply_to_corrected=True,
    )
    render_reviewed_docx(parsed, [action], reviewed)
    evidence = _evidence_for_actions(source, reviewed, [action])
    comparison = compare_docx_documents(source, reviewed)

    assert any(item.code == "numbering_not_projected" for item in comparison.diagnostics)
    provenance = _public_api("attribute_docx_changes")(
        comparison,
        action_source=source,
        review_evidence=evidence,
    )

    assert _one(provenance, "text").status == "verified_action"


def test_semantic_diagnostic_only_downgrades_its_matching_event_type(tmp_path: Path) -> None:
    source, reviewed, _corrected, action = _review_fixture(tmp_path)
    evidence = _evidence(source, reviewed, [action])
    comparison = compare_docx_documents(source, reviewed)
    text_change = comparison.text_changes[0]
    comparison = replace(
        comparison,
        diagnostics=(
            *comparison.diagnostics,
            ComparisonDiagnostic(
                code="text_diff_workload_fallback",
                message="range could not be isolated reliably",
                side="both",
                locator=text_change.pair_id,
            ),
        ),
    )

    provenance = _public_api("attribute_docx_changes")(
        comparison,
        action_source=source,
        review_evidence=evidence,
    )

    assert _one(provenance, "text").status == "unknown"
    assert {event.status for event in _events(provenance, "revision")} == {"verified_action"}


def test_blocked_and_comment_only_actions_do_not_claim_coincident_text_edits(
    tmp_path: Path,
) -> None:
    source = tmp_path / "blocked-input.docx"
    reviewed = tmp_path / "blocked-reviewed.docx"
    corrected = tmp_path / "blocked-corrected.docx"
    document = Document()
    document.add_paragraph("Old clause.")
    document.add_paragraph("Old note.")
    document.save(source)

    parsed = load_docx(source)
    paragraphs = tuple(parsed.iter_paragraphs())
    blocked_action = ReviewAction(
        id="blocked-replacement",
        scope=ReviewScope.PARAGRAPH,
        action_type=ReviewActionType.REPLACE_TEXT,
        node_id=paragraphs[0].id,
        original_text="Old",
        replacement_text="New",
        comment="Blocked from corrected output.",
        status=ActionStatus.APPLIED,
        apply_to_corrected=True,
        metadata={"blocked_from_corrected": True},
    )
    comment_only_action = ReviewAction(
        id="comment-only",
        scope=ReviewScope.PARAGRAPH,
        action_type=ReviewActionType.COMMENT,
        node_id=paragraphs[1].id,
        original_text="Old",
        replacement_text="New",
        comment="This only creates a comment.",
        status=ActionStatus.APPLIED,
        apply_to_corrected=True,
    )
    actions = [blocked_action, comment_only_action]
    render_reviewed_docx(parsed, actions, reviewed)
    evidence = _evidence_for_actions(source, reviewed, actions)

    reviewed_result = _public_api("attribute_docx_changes")(
        compare_docx_documents(source, reviewed),
        action_source=source,
        review_evidence=evidence,
    )
    action_comments = _events(reviewed_result, "comment")
    assert len(action_comments) == 2
    assert {event.status for event in action_comments} == {"verified_action"}
    assert {event.action_ids for event in action_comments} == {
        ("blocked-replacement",),
        ("comment-only",),
    }
    assert _events(reviewed_result, "revision") == []

    corrected_document = Document(source)
    corrected_paragraphs = corrected_document.paragraphs
    corrected_paragraphs[0].runs[0].text = "New clause."
    corrected_paragraphs[1].runs[0].text = "New note."
    corrected_document.save(corrected)
    corrected_result = _public_api("attribute_docx_changes")(
        compare_docx_documents(source, corrected),
        action_source=source,
        review_evidence=evidence,
        transition_evidence=_transition(reviewed, corrected),
    )

    text_changes = _events(corrected_result, "text")
    assert len(text_changes) == 2
    assert {event.status for event in text_changes} == {"unknown"}
    assert all(event.action_ids == () for event in text_changes)


def _replace_comment_text(comment: Comment, text: str) -> None:
    paragraphs = comment.paragraphs
    lines = text.splitlines() or [""]
    paragraphs[0].text = lines[0]
    for paragraph in paragraphs[1:]:
        paragraph._element.getparent().remove(paragraph._element)
    for line in lines[1:]:
        comment.add_paragraph(line)


def _historical_comment_fixture(
    tmp_path: Path,
    *,
    with_lineage: bool,
    comment_text: str = _HISTORICAL_ACTION_COMMENT,
) -> tuple[Path, Path, Path, ReviewAction, Any]:
    source = tmp_path / "input.docx"
    reviewed = tmp_path / "reviewed.docx"
    corrected = tmp_path / "corrected.docx"

    document = Document()
    paragraph = document.add_paragraph("The old marker is alpha.")
    document.add_comment(runs=paragraph.runs, text="Source note.", author="Reviewer")
    document.save(source)

    parsed = load_docx(source)
    target = next(parsed.iter_paragraphs())
    action_values: dict[str, object] = {
        "id": "replace-marker",
        "scope": ReviewScope.PARAGRAPH,
        "action_type": ReviewActionType.REPLACE_TEXT,
        "node_id": target.id,
        "original_text": "alpha",
        "replacement_text": "beta",
        "comment": "Replace marker.",
        "status": ActionStatus.APPLIED,
        "apply_to_corrected": True,
        "category": "term",
        "policy_reason": "policy check",
        "evidence_refs": [EvidenceRef(locator="body:p:0")],
        "references": [ReviewReference(source="source-ref", label="Section 1")],
    }
    if with_lineage:
        action_values["lineage"] = (FindingLineageEvent(kind="review", node_id=target.id),)
    action = ReviewAction(**action_values)
    assert format_legacy_action_comment(action) == _HISTORICAL_ACTION_COMMENT
    render_reviewed_docx(parsed, [action], reviewed, comment_author="Reviewer")
    accept_all_revisions(reviewed, corrected)

    reviewed_document = Document(reviewed)
    action_comment = next(
        comment
        for comment in reviewed_document.comments
        if comment.text == format_action_comment(action)
    )
    _replace_comment_text(action_comment, comment_text)
    reviewed_document.save(reviewed)

    evidence = (
        _evidence_for_actions(source, reviewed, [action])
        if with_lineage
        else _legacy_evidence_for_actions(source, reviewed, action)
    )
    return source, reviewed, corrected, action, evidence


def _review_fixture(tmp_path: Path) -> tuple[Path, Path, Path, ReviewAction]:
    source = tmp_path / "input.docx"
    reviewed = tmp_path / "reviewed.docx"
    corrected = tmp_path / "corrected.docx"

    document = Document()
    paragraph = document.add_paragraph("Termin zapłaty wynosi 14 dni od otrzymania faktury.")
    document.add_comment(runs=paragraph.runs, text=_SOURCE_COMMENT, author="Dike")
    document.save(source)

    parsed = load_docx(source)
    target = next(item for item in parsed.iter_paragraphs() if "14 dni" in item.text)
    action = ReviewAction(
        id=_ACTION_ID,
        scope=ReviewScope.PARAGRAPH,
        action_type=ReviewActionType.REPLACE_TEXT,
        node_id=target.id,
        original_text="14 dni",
        replacement_text="30 dni",
        comment=_ACTION_COMMENT,
        status=ActionStatus.APPLIED,
        apply_to_corrected=True,
        confidence=0.99,
    )
    render_reviewed_docx(parsed, [action], reviewed, comment_author="Dike")
    accept_all_revisions(reviewed, corrected)
    return source, reviewed, corrected, action


def _evidence(source: Path, reviewed: Path, actions: list[ReviewAction]) -> Any:
    evidence_type = _public_api("ReviewActionEvidence")
    # This hash is the independently recorded canonical digest for this complete
    # fixture action, rather than a value derived from the implementation under test.
    assert _action_digest(actions) == _ACTION_SHA256
    return evidence_type(
        source_sha256=_sha256(source),
        reviewed_sha256=_sha256(reviewed),
        actions_sha256=_ACTION_SHA256,
        source_revision_signatures=(),
        actions=tuple(actions),
    )


def _evidence_for_actions(source: Path, reviewed: Path, actions: list[ReviewAction]) -> Any:
    evidence_type = _public_api("ReviewActionEvidence")
    return evidence_type(
        source_sha256=_sha256(source),
        reviewed_sha256=_sha256(reviewed),
        actions_sha256=_action_digest(actions),
        source_revision_signatures=reviewkit.revision_signatures(source),
        actions=tuple(actions),
    )


def _legacy_evidence_for_actions(source: Path, reviewed: Path, action: ReviewAction) -> Any:
    legacy_payload = action.model_dump(mode="json", by_alias=True)
    legacy_payload.pop("lineage")
    legacy_action = ReviewAction.model_validate(legacy_payload)
    evidence_type = _public_api("ReviewActionEvidence")
    return evidence_type(
        source_sha256=_sha256(source),
        reviewed_sha256=_sha256(reviewed),
        actions_sha256=_json_digest([legacy_payload]),
        source_revision_signatures=reviewkit.revision_signatures(source),
        actions=(legacy_action,),
    )


def _transition(reviewed: Path, corrected: Path) -> Any:
    transition_type = _public_api("DocumentTransitionEvidence")
    return transition_type(input_sha256=_sha256(reviewed), output_sha256=_sha256(corrected))


def _action_digest(actions: list[ReviewAction]) -> str:
    canonical = reviewkit.canonical_action_dump(actions)
    encoded = json.dumps(canonical, sort_keys=True, ensure_ascii=False, separators=(",", ":"))
    return hashlib.sha256(encoded.encode("utf-8")).hexdigest()


def _json_digest(payload: object) -> str:
    encoded = json.dumps(payload, sort_keys=True, ensure_ascii=False, separators=(",", ":"))
    return hashlib.sha256(encoded.encode("utf-8")).hexdigest()


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _public_api(name: str) -> Any:
    value = getattr(reviewkit, name, None)
    assert value is not None, f"ReviewKit must expose {name} for verified DOCX attribution"
    return value


def _events(provenance: Any, event_type: str) -> list[Any]:
    return [event for event in provenance.changes if event.event_type == event_type]


def _one(provenance: Any, event_type: str) -> Any:
    events = _events(provenance, event_type)
    assert len(events) == 1
    return events[0]


def _by_change_id(provenance: Any, change_id: str) -> Any:
    return next(event for event in provenance.changes if event.change_id == change_id)
